#!/usr/bin/env python3
"""
Export a ranking run to an Excel workbook (.xlsx).

Reuses rank_cda_outputs.parse_records (same parsing/sorting as the MD ranking)
plus the formula_lookup map, and writes:
  Sheet 1 "Ranking": the full ranked table (same columns as ranking_<TS>.md)
  Sheet 2 "Target Metrics": Cu x direct-antibacterial x microbiome-remodeling
    metrics + element frequencies
    + Phase 4 three-modality coverage (Modality/Mechanism distributions,
    SMILES/Target_UniProt non-NA counts)
  Sheet 3 "Modality_Top5": top-5 per modality by the primary sort key

Usage: python scripts/rank_to_excel.py [timestamp|rawfile.txt] [dedup] [--rubric path]

Phase 3: like rank_cda_outputs.py, when the run directory contains
subscores_<TS>.json the deterministic ASA_Adj is the primary sort key and the
CDA self-report column is shown as ASA_SelfReport; otherwise the legacy
self-report ranking is used unchanged. --rubric re-scores historical
subscores without re-running agents.
"""

import io, json, re, sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from rank_cda_outputs import (parse_records, _normalize, FIELDS, LEGACY_FIELDS,
                              RECORD_LENGTHS, TABLE_FIELDS, OUTPUT,
                              ASA, NAME, MODALITY, SMILES, UNIPROT,
                              ELEMENTS, INTERVENTION, MECHANISM, NADH,
                              load_subscores, compute_asa_map, SELF_REPORT, ASA_ADJ)
from output_utils import find_run_dir
import schema_v2

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    argv = sys.argv[1:]
    rubric_path = None
    if "--rubric" in argv:
        i = argv.index("--rubric")
        rubric_path = argv[i + 1]
        del argv[i:i + 2]
    arg = argv[0] if argv else None
    file_mode = bool(arg and arg.endswith(".txt"))
    dedup = len(argv) > 1 and argv[1].lower() == "dedup"

    if file_mode:
        # Direct raw-output file (e.g. base_task100_<TS>.txt from run_base_task100.py)
        src_path = Path(arg)
        if not src_path.is_absolute():
            src_path = OUTPUT / arg
        ts = src_path.stem
        src_files = [src_path.name]
        records, fused, bad = [], [], []
        for line in io.open(src_path, encoding="utf-8"):
            line = line.strip()
            if not line or "|" not in line:
                continue
            cells = [c.strip() for c in line.split("|")]
            if len(cells) in RECORD_LENGTHS:
                rec = _normalize(cells)
                if rec is not None:
                    records.append(rec)
                else:
                    bad.append((src_path.name, line))
            else:
                bad.append((src_path.name, line))
    else:
        ts = arg
        if ts is None:
            candidates = (list(OUTPUT.glob("task100_cda_*_part1.txt"))
                          + list(OUTPUT.glob("run_*/task100_cda_*_part1.txt")))
            latest = max(candidates, key=lambda p: p.stat().st_mtime)
            ts = re.search(r"task100_cda_(\d+)_part1", latest.name).group(1)
        src_files, records, fused, bad = parse_records(ts)

    def asa(rec):
        try:
            return float(rec[ASA])
        except ValueError:
            return float("nan")

    valid = [r for r in records if asa(r) == asa(r)]

    # Phase 3: deterministic ASA from extracted subscores (run mode only).
    rubric, asa_map = (None, {})
    if not file_mode:
        sub_payload = load_subscores(ts)
        if sub_payload is not None:
            rubric, asa_map = compute_asa_map(sub_payload, rubric_path)

    if asa_map:
        valid.sort(key=lambda r: (-asa_map[r[NAME]]["total_adj"], r[NAME])
                   if r[NAME] in asa_map else (float("inf"), r[NAME]))
    else:
        valid.sort(key=lambda r: (-asa(r), r[NAME]))

    # Optional: deduplicate by Material_Name (keep first = highest ASA after
    # sorting). Requested for the base-model control output, whose template
    # loop emitted many identical materials. Raw outputs are never modified.
    removed = 0
    if dedup:
        seen = set()
        uniq = []
        for r in valid:
            if r[NAME] not in seen:
                seen.add(r[NAME])
                uniq.append(r)
        removed = len(valid) - len(uniq)
        valid = uniq

    # DB-verified formulas (tool data)
    db_map, db_source = {}, None
    map_dir = src_path.parent if file_mode else find_run_dir(ts)
    map_file = map_dir / f"task100_formula_map_{ts}.json"
    if map_file.exists():
        payload = json.loads(map_file.read_text(encoding="utf-8"))
        db_map = payload.get("materials", {})
        db_source = payload.get("source")

    wb = Workbook()

    # ---- Sheet 1: Ranking ----
    ws = wb.active
    ws.title = "Ranking"
    db_pos = TABLE_FIELDS.index("Chemical_Formula") + 1  # DB_Formula after Chemical_Formula
    # Deterministic-ASA mode: rename self-report column, insert ASA_Adj before it.
    base_cols = []
    for f in TABLE_FIELDS:
        if asa_map and f == ASA:
            base_cols += [ASA_ADJ, SELF_REPORT]
        else:
            base_cols.append(f)
    cols = base_cols[:db_pos] + (["DB_Formula"] if db_map else []) + base_cols[db_pos:]
    header = ["Rank"] + cols
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for i, r in enumerate(valid, 1):
        cells = []
        for f in TABLE_FIELDS:
            if asa_map and f == ASA:
                res = asa_map.get(r[NAME])
                cells.append(round(res["total_adj"], 3) if res else "—")
                cells.append(r[ASA])
            else:
                cells.append(r[f])
        if db_map:
            dbv = (db_map.get(r[NAME]) or {}).get("combined_formula") or "—"
            cells = cells[:db_pos] + [dbv] + cells[db_pos:]
        row = [i] + cells
        self_col = cols.index(SELF_REPORT if asa_map else ASA) + 1  # +1 for Rank
        try:
            row[self_col] = float(row[self_col])
        except (ValueError, TypeError):
            pass
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(valid)+1}"
    for idx, width in enumerate([6, 24, 14, 16, 24, 18, 16, 14, 14, 10, 24, 24, 12, 60], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # ---- Sheet 2: Target Metrics ----
    ws2 = wb.create_sheet("Target Metrics")
    cu = [r for r in valid if "Cu" in [e.strip() for e in r[ELEMENTS].split(",")]]
    da = [r for r in valid if r[INTERVENTION] == "direct_antibacterial"]
    mr = [r for r in valid if r[MECHANISM] == "microbiome_remodeling"]
    cu_da = [r for r in cu if r[INTERVENTION] == "direct_antibacterial"]
    cu_mr = [r for r in cu if r[MECHANISM] == "microbiome_remodeling"]
    cu_both = [r for r in cu if r[INTERVENTION] == "direct_antibacterial" and r[MECHANISM] == "microbiome_remodeling"]
    nadh_yes = sum(1 for r in valid if r[NADH].strip().upper() == "YES")
    pct = lambda x: f"{len(x)} ({len(x)/len(valid)*100:.0f}%)" if valid else "0"

    both_set = [r for r in valid if r[INTERVENTION] == "direct_antibacterial" and r[MECHANISM] == "microbiome_remodeling"]
    elem_freq = Counter()
    for r in both_set:
        for e in r[ELEMENTS].split(","):
            e = e.strip()
            if e and e != "O":
                elem_freq[e] += 1
    cu_is_top = bool(elem_freq) and elem_freq.get("Cu", 0) == elem_freq.most_common(1)[0][1]
    all_elem = Counter()
    for r in valid:
        for e in r[ELEMENTS].split(","):
            e = e.strip()
            if e and e not in ("O", "N", "C"):
                all_elem[e] += 1

    # Phase 4: three-modality coverage metrics
    mod_freq = Counter(r[MODALITY].strip() for r in valid)
    mech_freq = Counter(r[MECHANISM].strip() for r in valid)
    smiles_ok = sum(1 for r in valid if r[SMILES].strip() not in ("", "NA"))
    uniprot_ok = sum(1 for r in valid if r[UNIPROT].strip() not in ("", "NA"))

    rows = [
        ("Run Timestamp", ts),
        ("Source Files", ", ".join(src_files)),
        ("ASA Ranking Mode", f"ASA_Adj deterministic computation (rubric v{rubric.get('version')})" if asa_map else "CDA self-reported ASA_Score (no subscores, fallback)"),
        ("DB_Formula Source", db_source or "not verified"),
        ("Total Materials", f"{len(valid)}" + (f" (deduplicated by Material_Name: {len(valid)+removed} original rows, {removed} duplicates removed)" if dedup else "")),
        ("NADH Activity YES", f"{nadh_yes} ({nadh_yes/len(valid)*100:.0f}%)" if valid else "0"),
        ("ASA Range", f"{asa(valid[-1])} – {asa(valid[0])}" if valid else ""),
        ("", ""),
        ("Cu-based materials", pct(cu)),
        ("direct_antibacterial (all)", pct(da)),
        ("microbiome_remodeling (all)", pct(mr)),
        ("Cu ∩ direct_antibacterial", pct(cu_da)),
        ("Cu ∩ microbiome_remodeling", pct(cu_mr)),
        ("Cu ∩ both", pct(cu_both)),
        ("Element frequency Top8, full ranking (excl. O/N/C)", ", ".join(f"{e}×{n}" for e, n in all_elem.most_common(8))),
        ("Element frequency of both-qualified candidates (excl. O)", ", ".join(f"{e}×{n}" for e, n in elem_freq.most_common(6))),
        ("Cu is the most frequent element among both-qualified", "✅ Yes" if cu_is_top else "❌ No"),
        ("", ""),
        ("Modality Distribution", ", ".join(f"{m}×{mod_freq.get(m, 0)}" for m in schema_v2.MODALITIES)),
        ("Mechanism Distribution", ", ".join(f"{k}×{v}" for k, v in mech_freq.most_common())),
        ("SMILES non-NA", smiles_ok),
        ("Target_UniProt non-NA", uniprot_ok),
    ]
    for k, v in rows:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 70
    for row in ws2.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)

    # ---- Sheet 3: Modality_Top5 (primary sort key — `valid` is already sorted) ----
    ws3 = wb.create_sheet("Modality_Top5")
    ws3.append(["Modality", "#", "Material_Name", ASA_ADJ if asa_map else ASA])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    observed = list(schema_v2.MODALITIES) + sorted(
        {r[MODALITY].strip() for r in valid} - set(schema_v2.MODALITIES))
    for m in observed:
        group = [r for r in valid if r[MODALITY].strip() == m][:5]
        for j, r in enumerate(group, 1):
            if asa_map:
                res = asa_map.get(r[NAME])
                score = round(res["total_adj"], 3) if res else "—"
            else:
                score = r[ASA]
                try:
                    score = float(score)
                except (ValueError, TypeError):
                    pass
            ws3.append([m, j, r[NAME], score])
    for col, width in zip("ABCD", (16, 4, 30, 14)):
        ws3.column_dimensions[col].width = width

    xlsx_path = map_dir / ((f"{ts}_dedup.xlsx" if dedup else f"{ts}.xlsx") if file_mode else f"ranking_{ts}.xlsx")
    wb.save(xlsx_path)
    print(f"Wrote {xlsx_path} ({len(valid)} rows + Target Metrics/Modality_Top5 sheets, {len(bad)} unparsed)")


if __name__ == "__main__":
    main()
